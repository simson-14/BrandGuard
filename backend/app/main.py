import json
import re
from io import BytesIO

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from .analysis import analyze_content
from .database import connect, document_dict, initialize_database, row_dict


app = FastAPI(title="BrandGuard API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_methods=["*"],
    allow_headers=["*"],
)

RULE_TYPES = {"banned", "preferred", "tone", "claim", "fact", "required"}
SEVERITIES = {"High", "Medium", "Low"}


class BrandInput(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    description: str = Field(default="", max_length=2000)


class RuleInput(BaseModel):
    type: str
    term: str = Field(min_length=1, max_length=500)
    description: str = Field(min_length=1, max_length=1000)
    replacement: str = Field(default="", max_length=500)
    severity: str = "Medium"


class ReviewInput(BaseModel):
    status: str | None = None
    reviewer_comments: str | None = Field(default=None, max_length=5000)


class DecisionInput(BaseModel):
    decision: str


def validate_rule(payload):
    if payload.type not in RULE_TYPES:
        raise HTTPException(422, "Unsupported rule type.")
    if payload.severity not in SEVERITIES:
        raise HTTPException(422, "Severity must be High, Medium, or Low.")


def require_document(db, document_id):
    row = db.execute("SELECT * FROM documents WHERE id = ?", (document_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Document not found.")
    return row


@app.on_event("startup")
def startup():
    initialize_database()


@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.get("/api/brands")
def get_brands():
    with connect() as db:
        return [dict(row) for row in db.execute("SELECT * FROM brands ORDER BY id")]


@app.post("/api/brands")
def create_brand(payload: BrandInput):
    with connect() as db:
        cursor = db.execute(
            "INSERT INTO brands(name, description) VALUES (?, ?)",
            (payload.name.strip(), payload.description.strip()),
        )
        return row_dict(db.execute("SELECT * FROM brands WHERE id = ?", (cursor.lastrowid,)).fetchone())


@app.put("/api/brands/{brand_id}")
def update_brand(brand_id: int, payload: BrandInput):
    with connect() as db:
        cursor = db.execute(
            "UPDATE brands SET name = ?, description = ? WHERE id = ?",
            (payload.name.strip(), payload.description.strip(), brand_id),
        )
        if not cursor.rowcount:
            raise HTTPException(404, "Brand not found.")
        return row_dict(db.execute("SELECT * FROM brands WHERE id = ?", (brand_id,)).fetchone())


@app.get("/api/brands/{brand_id}/rules")
def get_rules(brand_id: int):
    with connect() as db:
        if not db.execute("SELECT 1 FROM brands WHERE id = ?", (brand_id,)).fetchone():
            raise HTTPException(404, "Brand not found.")
        return [dict(row) for row in db.execute("SELECT * FROM rules WHERE brand_id = ? ORDER BY id", (brand_id,))]


@app.post("/api/brands/{brand_id}/rules")
def create_rule(brand_id: int, payload: RuleInput):
    validate_rule(payload)
    with connect() as db:
        if not db.execute("SELECT 1 FROM brands WHERE id = ?", (brand_id,)).fetchone():
            raise HTTPException(404, "Brand not found.")
        cursor = db.execute(
            "INSERT INTO rules(brand_id, type, term, description, replacement, severity) VALUES (?, ?, ?, ?, ?, ?)",
            (brand_id, payload.type, payload.term.strip(), payload.description.strip(), payload.replacement.strip(), payload.severity),
        )
        return row_dict(db.execute("SELECT * FROM rules WHERE id = ?", (cursor.lastrowid,)).fetchone())


@app.put("/api/rules/{rule_id}")
def update_rule(rule_id: int, payload: RuleInput):
    validate_rule(payload)
    with connect() as db:
        cursor = db.execute(
            "UPDATE rules SET type = ?, term = ?, description = ?, replacement = ?, severity = ? WHERE id = ?",
            (payload.type, payload.term.strip(), payload.description.strip(), payload.replacement.strip(), payload.severity, rule_id),
        )
        if not cursor.rowcount:
            raise HTTPException(404, "Rule not found.")
        return row_dict(db.execute("SELECT * FROM rules WHERE id = ?", (rule_id,)).fetchone())


@app.delete("/api/rules/{rule_id}")
def delete_rule(rule_id: int):
    with connect() as db:
        cursor = db.execute("DELETE FROM rules WHERE id = ?", (rule_id,))
        if not cursor.rowcount:
            raise HTTPException(404, "Rule not found.")
        return {"deleted": True}


def extract_docx(filename, file_bytes):
    if not filename.lower().endswith(".docx"):
        raise HTTPException(415, "Only .docx files are supported.")
    if not file_bytes or len(file_bytes) > 10 * 1024 * 1024:
        raise HTTPException(400, "The DOCX file is empty or exceeds the 10 MB limit.")
    try:
        document = Document(BytesIO(file_bytes))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        paragraphs.extend(
            cell.text for table in document.tables for row in table.rows
            for cell in row.cells if cell.text.strip()
        )
    except Exception as error:
        raise HTTPException(400, "This file is not a valid DOCX document.") from error
    text = "\n".join(paragraphs).strip()
    if not text:
        raise HTTPException(400, "No readable text was found in this DOCX document.")
    return text


def analyze_and_save(db, title, content, content_type, brand_id):
    brand = db.execute("SELECT * FROM brands WHERE id = ?", (brand_id,)).fetchone()
    if not brand:
        raise HTTPException(404, "Brand not found.")
    rules = [dict(row) for row in db.execute("SELECT * FROM rules WHERE brand_id = ?", (brand_id,))]
    analysis = analyze_content(content, rules)
    cursor = db.execute(
        "INSERT INTO documents(brand_id, title, content_type, content, analysis) VALUES (?, ?, ?, ?, ?)",
        (brand_id, title.strip(), content_type, content, json.dumps(analysis)),
    )
    return document_dict(db.execute("SELECT * FROM documents WHERE id = ?", (cursor.lastrowid,)).fetchone())


@app.post("/api/documents")
def submit_document(payload: dict):
    title = str(payload.get("title", "")).strip()
    content = str(payload.get("content", "")).strip()
    content_type = str(payload.get("content_type", "website"))
    try:
        brand_id = int(payload.get("brand_id", 1))
    except (ValueError, TypeError) as error:
        raise HTTPException(422, "A valid brand is required.") from error
    if not title or len(title) > 200:
        raise HTTPException(422, "Title is required and must be 200 characters or fewer.")
    if not content or len(content) > 200_000:
        raise HTTPException(422, "Content is required and must be 200,000 characters or fewer.")
    if content_type not in {"website", "email", "social"}:
        raise HTTPException(422, "Choose website, email, or social media as the content type.")
    with connect() as db:
        return analyze_and_save(db, title, content, content_type, brand_id)


@app.post("/api/documents/upload")
async def upload_document(
    title: str = Form(...),
    content_type: str = Form(...),
    brand_id: int = Form(1),
    file: UploadFile = File(...),
):
    if not title.strip() or len(title) > 200:
        raise HTTPException(422, "Title is required and must be 200 characters or fewer.")
    if content_type not in {"website", "email", "social"}:
        raise HTTPException(422, "Choose website, email, or social media as the content type.")
    content = extract_docx(file.filename or "", await file.read())
    if len(content) > 200_000:
        raise HTTPException(422, "Extracted text exceeds the 200,000 character limit.")
    with connect() as db:
        return analyze_and_save(db, title, content, content_type, brand_id)


@app.get("/api/documents")
def get_documents(status: str | None = None, search: str | None = None):
    query = "SELECT * FROM documents WHERE 1 = 1"
    params = []
    if status and status != "All statuses":
        query += " AND status = ?"
        params.append(status)
    if search:
        query += " AND (title LIKE ? OR content LIKE ?)"
        params.extend([f"%{search}%", f"%{search}%"])
    query += " ORDER BY submitted_at DESC, id DESC"
    with connect() as db:
        return [document_dict(row) for row in db.execute(query, params)]


@app.get("/api/documents/{document_id}")
def get_document(document_id: int):
    with connect() as db:
        return document_dict(require_document(db, document_id))


@app.post("/api/documents/{document_id}/issues/{issue_id}/decision")
def decide_suggestion(document_id: int, issue_id: str, payload: DecisionInput):
    if payload.decision not in {"accepted", "rejected"}:
        raise HTTPException(422, "Decision must be accepted or rejected.")
    with connect() as db:
        row = require_document(db, document_id)
        analysis = json.loads(row["analysis"])
        issue = next((item for item in analysis["issues"] if item["id"] == issue_id), None)
        if not issue:
            raise HTTPException(404, "Issue not found.")
        issue["suggestion_status"] = payload.decision
        db.execute("UPDATE documents SET analysis = ? WHERE id = ?", (json.dumps(analysis), document_id))
        db.execute(
            "INSERT INTO review_events(document_id, issue_id, action) VALUES (?, ?, ?)",
            (document_id, issue_id, payload.decision),
        )
        return document_dict(require_document(db, document_id))


@app.post("/api/documents/{document_id}/review")
def review_document(document_id: int, payload: ReviewInput):
    allowed = {"Awaiting Review", "Needs Changes", "Approved"}
    if payload.status is not None and payload.status not in allowed:
        raise HTTPException(422, "Choose Awaiting Review, Needs Changes, or Approved.")
    with connect() as db:
        row = require_document(db, document_id)
        status = payload.status or row["status"]
        comments = payload.reviewer_comments if payload.reviewer_comments is not None else row["reviewer_comments"]
        approval_at = row["approval_at"]
        if payload.status == "Approved":
            approval_at = __import__("datetime").datetime.now(__import__("datetime").timezone.utc).isoformat()
        elif payload.status is not None:
            approval_at = None
        db.execute(
            "UPDATE documents SET status = ?, reviewer_comments = ?, approval_at = ? WHERE id = ?",
            (status, comments, approval_at, document_id),
        )
        if payload.status is not None:
            db.execute(
                "INSERT INTO review_events(document_id, action, comment) VALUES (?, ?, ?)",
                (document_id, status, comments),
            )
        elif payload.reviewer_comments is not None:
            db.execute(
                "INSERT INTO review_events(document_id, action, comment) VALUES (?, 'Comment', ?)",
                (document_id, comments, ),
            )
        return document_dict(require_document(db, document_id))


@app.get("/api/dashboard")
def dashboard():
    with connect() as db:
        documents = [document_dict(row) for row in db.execute("SELECT * FROM documents")]
    issues = [issue for document in documents for issue in document["analysis"]["issues"]]
    categories = {}
    for issue in issues:
        categories[issue["category"]] = categories.get(issue["category"], 0) + 1
    return {
        "documents_checked": len(documents),
        "violations_detected": len(issues),
        "high_severity": sum(issue["severity"] == "High" for issue in issues),
        "awaiting_approval": sum(doc["status"] == "Awaiting Review" for doc in documents),
        "recent_documents": documents[:6],
        "category_counts": categories,
    }


@app.get("/api/documents/{document_id}/events")
def get_review_events(document_id: int):
    with connect() as db:
        require_document(db, document_id)
        return [dict(row) for row in db.execute(
            "SELECT * FROM review_events WHERE document_id = ? ORDER BY created_at DESC, id DESC",
            (document_id,),
        )]
